import streamlit as st
import pandas as pd
import json
import re
from io import BytesIO

from groq import Groq
from pypdf import PdfReader
from docx import Document


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="HireTech AI",
    page_icon="🤖",
    layout="wide"
)


# ============================================================
# DEFAULT JOB DESCRIPTION
# ============================================================

DEFAULT_JOB_DESCRIPTION = """
We are looking for a qualified Hydrologist to join our water resources and
engineering team.

The candidate should have a degree in Hydrology, Water Resources Engineering,
Civil Engineering, Environmental Engineering, or a closely related field,
with relevant professional experience in hydrological and water resources
projects.

The candidate will be responsible for hydrological analysis, rainfall-runoff
modeling, flood studies, watershed and catchment analysis, drainage studies,
water resources assessment, flood risk assessment, and preparation of
technical reports.

The role may involve hydrologic and hydraulic modeling, analysis of rainfall
and streamflow data, flood frequency analysis, watershed delineation,
catchment modeling, stormwater analysis, river basin studies, water balance
analysis, and preparation of engineering and hydrology reports.

Experience with software and tools such as HEC-HMS, HEC-RAS, SWMM, ArcGIS,
QGIS, AutoCAD, Microsoft Excel, GIS, remote sensing tools, and other
hydrological or hydraulic modeling software will be considered highly
relevant.

The candidate should have strong knowledge of hydrology, hydraulics,
watershed management, precipitation analysis, rainfall-runoff processes,
flood modeling, drainage systems, water resources management, and
hydrological data analysis.

Strong analytical, numerical, problem-solving, report-writing,
communication, teamwork, and project coordination skills are required.

Experience with flood studies, river basin studies, watershed modeling,
stormwater management, groundwater studies, climate and rainfall analysis,
or related water resources projects will be considered valuable.
"""


# ============================================================
# DEFAULT KEYWORDS
# ============================================================

DEFAULT_KEYWORDS = (
    "Hydrology, Hydrologist, Water Resources, HEC-HMS, HEC-RAS, "
    "SWMM, ArcGIS, QGIS, GIS, Hydrologic Modeling, Hydraulic Modeling, "
    "Rainfall-Runoff Modeling, Watershed, Catchment, Flood Modeling, "
    "Flood Frequency Analysis, Flood Risk Assessment, Drainage, "
    "Stormwater, River Basin, Precipitation, Streamflow, SCS-CN, "
    "Unit Hydrograph, IDF Curves, Water Balance, Groundwater, "
    "Watershed Management, Water Resources Management"
)


# ============================================================
# SESSION STATE
# ============================================================

defaults = {
    "candidates": [],
    "analysis_errors": [],
    "job_title": "Hydrologist",
    "min_experience": 0,
    "required_keywords": DEFAULT_KEYWORDS,
    "job_description": DEFAULT_JOB_DESCRIPTION,
    "weight_experience": 30,
    "weight_technical": 25,
    "weight_education": 15,
    "weight_requirements": 20,
    "weight_certification": 10,
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


# ============================================================
# GROQ CONFIGURATION
# ============================================================

def get_groq_api_key():
    try:
        return str(st.secrets["GROQ_API_KEY"]).strip()
    except Exception:
        return ""


api_key = get_groq_api_key()

# Internal model — recruiter does not see this
GROQ_MODEL = "openai/gpt-oss-20b"


# ============================================================
# LEFT SIDEBAR
# ============================================================

with st.sidebar:

    st.title("⚙️ HireTech AI")

    st.markdown("### 📖 Instructions")

    st.markdown("""
    **1. Select the role**
    
    The application starts with a sample Hydrologist vacancy.
    You can change the job title for any recruitment position.

    **2. Review keywords**
    
    Default keywords are provided according to the sample vacancy.
    Add, remove, or change keywords as required.

    **3. Review the job description**
    
    A complete job description is provided by default.
    You can edit it for a specific vacancy.

    **4. Set minimum experience**
    
    Enter the minimum relevant experience required for the role.

    **5. Set scoring weightage**
    
    Adjust the importance of:
    - Experience
    - Technical Skills
    - Education
    - Job Requirements
    - Certifications

    **6. Upload CVs**
    
    Upload multiple PDF or DOCX CVs.

    **7. Analyze candidates**
    
    HireTech AI extracts candidate information, checks the CVs against
    the requirements, identifies matched and missing keywords, and
    calculates candidate scores.

    **8. Review results**
    
    Compare candidates, inspect individual profiles, review score
    breakdowns, and export the results to Excel.
    """)

    st.divider()

    st.subheader("⚖️ Scoring Weightage")

    st.caption(
        "Adjust the importance of each factor. "
        "The total does not have to equal 100%; scores are normalized automatically."
    )

    st.slider(
        "Experience",
        min_value=0,
        max_value=100,
        value=30,
        key="weight_experience"
    )

    st.slider(
        "Technical Skills",
        min_value=0,
        max_value=100,
        value=25,
        key="weight_technical"
    )

    st.slider(
        "Education",
        min_value=0,
        max_value=100,
        value=15,
        key="weight_education"
    )

    st.slider(
        "Job Requirements",
        min_value=0,
        max_value=100,
        value=20,
        key="weight_requirements"
    )

    st.slider(
        "Certifications",
        min_value=0,
        max_value=100,
        value=10,
        key="weight_certification"
    )

    total_weight = (
        st.session_state.weight_experience
        + st.session_state.weight_technical
        + st.session_state.weight_education
        + st.session_state.weight_requirements
        + st.session_state.weight_certification
    )

    st.metric(
        "Total Weight",
        f"{total_weight}%"
    )

    if total_weight == 100:
        st.success("Weightage totals 100%")
    else:
        st.info(
            "Weights are automatically normalized when calculating the final score."
        )


# ============================================================
# MAIN HEADER
# ============================================================

st.title("🤖 HireTech AI")
st.subheader("AI-Powered Recruitment & Candidate Ranking")

st.caption(
    "Analyze CVs against job requirements, compare candidates, "
    "identify strengths and gaps, and generate a ranked recruitment report."
)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def safe_string(value):

    if value is None:
        return ""

    if isinstance(value, str):
        return value.strip()

    if isinstance(value, dict):

        parts = []

        for key, val in value.items():

            if val is not None and str(val).strip():
                parts.append(f"{key}: {val}")

        return ", ".join(parts)

    if isinstance(value, list):

        parts = []

        for item in value:

            text = safe_string(item)

            if text:
                parts.append(text)

        return ", ".join(parts)

    return str(value).strip()


def safe_list(value):

    if value is None:
        return []

    if isinstance(value, list):

        result = []

        for item in value:

            text = safe_string(item)

            if text:
                result.append(text)

        return result

    if isinstance(value, dict):
        return [safe_string(value)]

    if isinstance(value, str):

        return [value.strip()] if value.strip() else []

    return [str(value)]


def safe_join(value, separator=", "):
    return separator.join(safe_list(value))


# ============================================================
# FILE TEXT EXTRACTION
# ============================================================

def extract_pdf_text(uploaded_file):

    uploaded_file.seek(0)

    reader = PdfReader(uploaded_file)

    pages = []

    for page in reader.pages:

        try:

            text = page.extract_text()

            if text:
                pages.append(text)

        except Exception:
            continue

    return "\n".join(pages)


def extract_docx_text(uploaded_file):

    uploaded_file.seek(0)

    document = Document(uploaded_file)

    paragraphs = []

    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                if cell.text.strip():
                    row_text.append(cell.text.strip())

            if row_text:
                paragraphs.append(" | ".join(row_text))

    return "\n".join(paragraphs)


def extract_text(uploaded_file):

    file_name = uploaded_file.name.lower()

    if file_name.endswith(".pdf"):
        return extract_pdf_text(uploaded_file)

    if file_name.endswith(".docx"):
        return extract_docx_text(uploaded_file)

    return ""


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(text):

    text = str(text).lower()

    text = text.replace("–", "-")
    text = text.replace("—", "-")
    text = text.replace("’", "'")

    text = re.sub(r"\s+", " ", text)

    return text.strip()


def normalize_keyword(keyword):

    keyword = normalize_text(keyword)

    keyword = re.sub(
        r"\s+",
        " ",
        keyword
    )

    return keyword.strip()


# ============================================================
# KEYWORD VARIANTS
# ============================================================

def get_keyword_variants(keyword):

    keyword = normalize_keyword(keyword)

    variants = {keyword}

    replacements = {

        "hec hms": [
            "hec-hms",
            "hec hms",
            "hec-hms software"
        ],

        "hec-hms": [
            "hec-hms",
            "hec hms"
        ],

        "hec ras": [
            "hec-ras",
            "hec ras",
            "hec-ras software"
        ],

        "hec-ras": [
            "hec-ras",
            "hec ras"
        ],

        "arc gis": [
            "arcgis",
            "arc gis"
        ],

        "arcgis": [
            "arcgis",
            "arc gis"
        ],

        "q gis": [
            "qgis",
            "q gis"
        ],

        "qgis": [
            "qgis",
            "q gis"
        ],

        "rainfall runoff modeling": [
            "rainfall-runoff modeling",
            "rainfall runoff modeling",
            "rainfall-runoff modelling",
            "rainfall runoff modelling"
        ],

        "hydrologic modeling": [
            "hydrologic modeling",
            "hydrologic modelling"
        ],

        "hydraulic modeling": [
            "hydraulic modeling",
            "hydraulic modelling"
        ],

        "stormwater": [
            "stormwater",
            "storm water"
        ],

        "streamflow": [
            "streamflow",
            "stream flow"
        ]
    }

    if keyword in replacements:
        variants.update(replacements[keyword])

    return list(variants)


# ============================================================
# KEYWORD MATCHING
# ============================================================

def find_keywords_in_cv(cv_text, required_keywords):

    normalized_cv = normalize_text(cv_text)

    matched = []
    missing = []

    for keyword in required_keywords:

        keyword = normalize_keyword(keyword)

        if not keyword:
            continue

        variants = get_keyword_variants(keyword)

        found = False

        for variant in variants:

            if variant in normalized_cv:
                found = True
                break

        if found:
            matched.append(keyword)
        else:
            missing.append(keyword)

    return matched, missing


def keyword_score(matched, required_keywords):

    total = len(required_keywords)

    if total == 0:
        return 0

    return round(
        (len(matched) / total) * 100,
        2
    )


# ============================================================
# JSON CLEANING
# ============================================================

def clean_json_response(response_text):

    if not response_text:
        return None

    text = response_text.strip()

    text = re.sub(
        r"```json",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"```",
        "",
        text
    )

    text = text.strip()

    first_brace = text.find("{")
    last_brace = text.rfind("}")

    if first_brace != -1 and last_brace != -1:
        text = text[first_brace:last_brace + 1]

    text = re.sub(
        r"[\x00-\x08\x0B\x0C\x0E-\x1F]",
        "",
        text
    )

    try:
        return json.loads(text)

    except json.JSONDecodeError:

        repaired = text

        repaired = repaired.replace("\n", " ")
        repaired = repaired.replace("\t", " ")

        try:
            return json.loads(repaired)
        except Exception:
            return None


# ============================================================
# AI CV ANALYSIS
# ============================================================

def analyze_cv(cv_text, job_description):

    if not api_key:

        raise ValueError(
            "GROQ_API_KEY is not configured in Streamlit Secrets."
        )

    client = Groq(
        api_key=api_key
    )

    prompt = f"""
You are an expert HR recruitment assistant.

Analyze the candidate CV against the job vacancy below.

The application is designed for recruitment across different professional
roles. Analyze the candidate according to the actual job title, keywords,
and job description provided.

JOB TITLE:
{st.session_state.job_title}

JOB DESCRIPTION:
{job_description}

CANDIDATE CV:
{cv_text}

Return ONLY valid JSON.

Return exactly this structure:

{{
    "candidate_name": "",
    "education": [],
    "years_relevant_experience": 0,
    "previous_roles": [],
    "technical_skills": [],
    "software_tools": [],
    "certifications": [],
    "relevant_experience_evidence": [],
    "strengths": [],
    "potential_gaps": []
}}

Rules:

1. candidate_name:
Extract the candidate's full name.

2. education:
Return readable education entries.
Example:
"M.Sc. Civil Engineering - University of XYZ - 2023"

3. years_relevant_experience:
Estimate relevant professional experience for this job.
Return a number.

4. previous_roles:
List previous job titles and organizations.

5. technical_skills:
List technical skills relevant to the vacancy.

6. software_tools:
List relevant software and technical tools.

7. certifications:
List relevant certifications.

8. relevant_experience_evidence:
Provide short evidence from the CV demonstrating relevant experience.

9. strengths:
List strengths relevant to the vacancy.

10. potential_gaps:
List missing or weaker areas relevant to the vacancy.

Do not invent information that is not supported by the CV.
"""


    response = client.chat.completions.create(

        model=GROQ_MODEL,

        messages=[
            {
                "role": "system",
                "content":
                    "You are a professional HR recruitment "
                    "and CV analysis assistant."
            },
            {
                "role": "user",
                "content": prompt
            }
        ],

        temperature=0,

        max_tokens=4000
    )

    response_text = (
        response
        .choices[0]
        .message
        .content
    )

    data = clean_json_response(
        response_text
    )

    if data is None:

        raise ValueError(
            "The AI returned an invalid JSON response. "
            "Please try again."
        )

    return data


# ============================================================
# SCORING
# ============================================================

def experience_score(years, minimum_required):

    try:

        years = float(years)
        minimum_required = float(minimum_required)

    except Exception:

        return 0

    if minimum_required <= 0:

        if years <= 0:
            return 0

        return min(
            100,
            years * 20
        )

    score = (
        years / minimum_required
    ) * 100

    return round(
        min(100, score),
        2
    )


def education_score(education):

    education_text = normalize_text(
        safe_join(education)
    )

    if not education_text:
        return 0

    # Hydrology / Water Resources
    if (
        "hydrology" in education_text
        or
        "water resources" in education_text
    ):
        return 100

    # Related engineering
    if (
        "civil engineering" in education_text
        or
        "environmental engineering" in education_text
    ):
        return 90

    if "engineering" in education_text:
        return 80

    if (
        "geology" in education_text
        or
        "geography" in education_text
    ):
        return 70

    return 40


def certification_score(certifications):

    certifications = safe_list(certifications)

    if not certifications:
        return 0

    return 100


def calculate_scores(
    candidate,
    required_keywords,
    min_experience,
    weights
):

    matched = candidate.get(
        "matched_keywords",
        []
    )

    years = candidate.get(
        "years_relevant_experience",
        0
    )

    education = candidate.get(
        "education",
        []
    )

    certifications = candidate.get(
        "certifications",
        []
    )

    requirement_score = keyword_score(
        matched,
        required_keywords
    )

    technical_score = requirement_score

    exp_score = experience_score(
        years,
        min_experience
    )

    edu_score = education_score(
        education
    )

    cert_score = certification_score(
        certifications
    )

    total_weight = sum(
        weights.values()
    )

    if total_weight <= 0:
        total_weight = 100

    final_score = (

        exp_score *
        weights["experience"]

        +

        technical_score *
        weights["technical"]

        +

        edu_score *
        weights["education"]

        +

        requirement_score *
        weights["requirements"]

        +

        cert_score *
        weights["certification"]

    ) / total_weight

    candidate["experience_score"] = round(
        exp_score,
        2
    )

    candidate["technical_score"] = round(
        technical_score,
        2
    )

    candidate["education_score"] = round(
        edu_score,
        2
    )

    candidate["requirements_score"] = round(
        requirement_score,
        2
    )

    candidate["certification_score"] = round(
        cert_score,
        2
    )

    candidate["final_score"] = round(
        final_score,
        2
    )

    return candidate


def recalculate_results():

    required_keywords = [

        normalize_keyword(x)

        for x in
        st.session_state.required_keywords.split(",")

        if normalize_keyword(x)
    ]

    weights = {

        "experience":
            st.session_state.weight_experience,

        "technical":
            st.session_state.weight_technical,

        "education":
            st.session_state.weight_education,

        "requirements":
            st.session_state.weight_requirements,

        "certification":
            st.session_state.weight_certification
    }

    for candidate in st.session_state.candidates:

        matched, missing = find_keywords_in_cv(

            candidate.get(
                "cv_text",
                ""
            ),

            required_keywords
        )

        candidate["matched_keywords"] = matched
        candidate["missing_keywords"] = missing

        calculate_scores(

            candidate,

            required_keywords,

            st.session_state.min_experience,

            weights
        )

    st.session_state.candidates.sort(

        key=lambda x:
            x.get(
                "final_score",
                0
            ),

        reverse=True
    )


# ============================================================
# JOB REQUIREMENTS
# ============================================================

st.header("📋 Job Requirements")

col1, col2 = st.columns(2)

with col1:

    st.text_input(
        "Job Title",
        key="job_title"
    )

    st.number_input(
        "Minimum Relevant Experience (Years)",
        min_value=0,
        max_value=50,
        step=1,
        key="min_experience"
    )


with col2:

    st.text_area(
        "Required Keywords",
        key="required_keywords",
        height=180,

        help=(
            "Default keywords are provided based on the sample vacancy. "
            "You can add, remove, or change them."
        )
    )

    st.caption(
        "💡 The default keywords can be changed for any job role."
    )


# ============================================================
# JOB DESCRIPTION
# ============================================================

st.subheader("📝 Job Description")

st.text_area(
    "Job Description",

    key="job_description",

    height=320,

    help=(
        "A sample job description is provided by default. "
        "Edit it according to the vacancy."
    )
)

st.caption(
    "💡 The default description can be customized for any position."
)


# ============================================================
# CV UPLOAD
# ============================================================

st.header("📄 Upload Candidate CVs")

uploaded_files = st.file_uploader(

    "Upload CVs in PDF or DOCX format",

    type=[
        "pdf",
        "docx"
    ],

    accept_multiple_files=True
)

if uploaded_files:

    st.info(
        f"📄 {len(uploaded_files)} CV(s) selected."
    )


# ============================================================
# ANALYZE BUTTON
# ============================================================

if st.button(

    "🚀 Analyze & Rank Candidates",

    type="primary",

    use_container_width=True
):

    st.session_state.candidates = []
    st.session_state.analysis_errors = []

    if not api_key:

        st.error(
            "GROQ_API_KEY is missing. "
            "Please add it to Streamlit Secrets."
        )

    elif not uploaded_files:

        st.warning(
            "Please upload at least one CV."
        )

    elif not st.session_state.job_description.strip():

        st.warning(
            "Please provide a job description."
        )

    else:

        progress = st.progress(0)

        total_files = len(
            uploaded_files
        )

        for index, uploaded_file in enumerate(
            uploaded_files
        ):

            try:

                cv_text = extract_text(
                    uploaded_file
                )

                if not cv_text.strip():

                    raise ValueError(
                        "Could not extract readable text from this CV."
                    )

                with st.spinner(
                    f"Analyzing {uploaded_file.name}..."
                ):

                    analysis = analyze_cv(
                        cv_text,
                        st.session_state.job_description
                    )

                candidate = {

                    "file_name":
                        uploaded_file.name,

                    "cv_text":
                        cv_text,

                    "candidate_name":
                        safe_string(
                            analysis.get(
                                "candidate_name",
                                ""
                            )
                        ),

                    "education":
                        safe_list(
                            analysis.get(
                                "education",
                                []
                            )
                        ),

                    "years_relevant_experience":
                        analysis.get(
                            "years_relevant_experience",
                            0
                        ),

                    "previous_roles":
                        safe_list(
                            analysis.get(
                                "previous_roles",
                                []
                            )
                        ),

                    "technical_skills":
                        safe_list(
                            analysis.get(
                                "technical_skills",
                                []
                            )
                        ),

                    "software_tools":
                        safe_list(
                            analysis.get(
                                "software_tools",
                                []
                            )
                        ),

                    "certifications":
                        safe_list(
                            analysis.get(
                                "certifications",
                                []
                            )
                        ),

                    "relevant_experience_evidence":
                        safe_list(
                            analysis.get(
                                "relevant_experience_evidence",
                                []
                            )
                        ),

                    "strengths":
                        safe_list(
                            analysis.get(
                                "strengths",
                                []
                            )
                        ),

                    "potential_gaps":
                        safe_list(
                            analysis.get(
                                "potential_gaps",
                                []
                            )
                        )
                }

                if not candidate["candidate_name"]:

                    candidate["candidate_name"] = (
                        uploaded_file.name
                    )

                st.session_state.candidates.append(
                    candidate
                )

            except Exception as e:

                st.session_state.analysis_errors.append(
                    f"{uploaded_file.name}: {str(e)}"
                )

            progress.progress(
                (index + 1) / total_files
            )

        if st.session_state.candidates:

            recalculate_results()

            st.success(
                f"✅ Successfully analyzed "
                f"{len(st.session_state.candidates)} candidate(s)."
            )


# ============================================================
# ERRORS
# ============================================================

if st.session_state.analysis_errors:

    st.subheader("⚠️ Analysis Errors")

    for error in st.session_state.analysis_errors:

        st.error(error)


# ============================================================
# RESULTS
# ============================================================

if st.session_state.candidates:

    st.header("🏆 Candidate Ranking")

    ranking_data = []

    for index, candidate in enumerate(
        st.session_state.candidates,
        start=1
    ):

        ranking_data.append({

            "Rank":
                index,

            "Candidate":
                candidate.get(
                    "candidate_name",
                    "Unknown"
                ),

            "Experience":
                candidate.get(
                    "years_relevant_experience",
                    0
                ),

            "Technical Score":
                candidate.get(
                    "technical_score",
                    0
                ),

            "Education Score":
                candidate.get(
                    "education_score",
                    0
                ),

            "Requirements Score":
                candidate.get(
                    "requirements_score",
                    0
                ),

            "Certification Score":
                candidate.get(
                    "certification_score",
                    0
                ),

            "Final Score":
                candidate.get(
                    "final_score",
                    0
                )
        })

    ranking_df = pd.DataFrame(
        ranking_data
    )

    st.dataframe(
        ranking_df,
        use_container_width=True,
        hide_index=True
    )


    # ========================================================
    # KEYWORD MATCHING
    # ========================================================

    st.header("🔎 Keyword Matching")

    required_keywords = [

        normalize_keyword(x)

        for x in
        st.session_state.required_keywords.split(",")

        if normalize_keyword(x)
    ]

    if required_keywords:

        keyword_rows = []

        for candidate in st.session_state.candidates:

            matched = candidate.get(
                "matched_keywords",
                []
            )

            missing = candidate.get(
                "missing_keywords",
                []
            )

            keyword_rows.append({

                "Candidate":
                    candidate.get(
                        "candidate_name",
                        "Unknown"
                    ),

                "Matched Keywords":
                    ", ".join(matched),

                "Missing Keywords":
                    ", ".join(missing),

                "Keyword Match %":
                    keyword_score(
                        matched,
                        required_keywords
                    )
            })

        keyword_df = pd.DataFrame(
            keyword_rows
        )

        st.dataframe(
            keyword_df,
            use_container_width=True,
            hide_index=True
        )

    else:

        st.info(
            "Enter keywords to see keyword matching."
        )


    # ========================================================
    # SCORE CHART
    # ========================================================

    st.header("📊 Candidate Scores")

    chart_df = (
        ranking_df[
            [
                "Candidate",
                "Final Score"
            ]
        ]
        .set_index("Candidate")
    )

    st.bar_chart(
        chart_df
    )


    # ========================================================
    # RECALCULATE
    # ========================================================

    if st.button(
        "🔄 Recalculate Ranking Using Current Weightage",
        use_container_width=True
    ):

        recalculate_results()

        st.success(
            "Ranking recalculated using the current weightage."
        )

        st.rerun()


    # ========================================================
    # CANDIDATE DETAILS
    # ========================================================

    st.header("👤 Candidate Details")

    candidate_options = [

        candidate.get(
            "candidate_name",
            "Unknown"
        )

        for candidate in
        st.session_state.candidates
    ]

    selected_detail_name = st.selectbox(

        "Select Candidate",

        candidate_options,

        key="candidate_detail_selector"
    )

    detail_candidate = next(

        (
            candidate

            for candidate in
            st.session_state.candidates

            if candidate.get(
                "candidate_name",
                "Unknown"
            ) == selected_detail_name

        ),

        None
    )

    if detail_candidate:

        st.subheader(
            detail_candidate.get(
                "candidate_name",
                "Unknown"
            )
        )

        col1, col2, col3, col4 = st.columns(4)

        with col1:

            st.metric(
                "Final Score",
                f"{detail_candidate.get('final_score', 0):.2f}%"
            )

        with col2:

            st.metric(
                "Experience",
                f"{detail_candidate.get('years_relevant_experience', 0)} years"
            )

        with col3:

            st.metric(
                "Keyword Match",
                f"{keyword_score(
                    detail_candidate.get(
                        'matched_keywords',
                        []
                    ),
                    required_keywords
                ):.2f}%"
            )

        with col4:

            st.metric(
                "Education Score",
                f"{detail_candidate.get('education_score', 0):.2f}%"
            )


        # ====================================================
        # MATCHED / MISSING
        # ====================================================

        st.subheader("🔎 Keyword Analysis")

        matched_keywords = detail_candidate.get(
            "matched_keywords",
            []
        )

        missing_keywords = detail_candidate.get(
            "missing_keywords",
            []
        )

        col1, col2 = st.columns(2)

        with col1:

            st.markdown("### ✅ Matched Keywords")

            if matched_keywords:

                for keyword in matched_keywords:
                    st.success(keyword)

            else:

                st.info(
                    "No required keywords matched."
                )

        with col2:

            st.markdown("### ❌ Missing Keywords")

            if missing_keywords:

                for keyword in missing_keywords:
                    st.error(keyword)

            else:

                st.success(
                    "All required keywords were found."
                )


        # ====================================================
        # EDUCATION
        # ====================================================

        st.subheader("🎓 Education")

        education = detail_candidate.get(
            "education",
            []
        )

        if education:

            for item in education:

                st.write(
                    f"• {safe_string(item)}"
                )

        else:

            st.write(
                "No education information found."
            )


        # ====================================================
        # PREVIOUS ROLES
        # ====================================================

        st.subheader("💼 Previous Roles")

        roles = detail_candidate.get(
            "previous_roles",
            []
        )

        if roles:

            for role in roles:

                st.write(
                    f"• {safe_string(role)}"
                )

        else:

            st.write(
                "No previous roles found."
            )


        # ====================================================
        # TECHNICAL SKILLS
        # ====================================================

        st.subheader("🛠️ Technical Skills")

        skills = detail_candidate.get(
            "technical_skills",
            []
        )

        if skills:

            st.write(
                ", ".join(
                    safe_list(skills)
                )
            )

        else:

            st.write(
                "No technical skills identified."
            )


        # ====================================================
        # SOFTWARE
        # ====================================================

        st.subheader("💻 Software & Tools")

        software = detail_candidate.get(
            "software_tools",
            []
        )

        if software:

            st.write(
                ", ".join(
                    safe_list(software)
                )
            )

        else:

            st.write(
                "No software tools identified."
            )


        # ====================================================
        # CERTIFICATIONS
        # ====================================================

        st.subheader("📜 Certifications")

        certifications = detail_candidate.get(
            "certifications",
            []
        )

        if certifications:

            for certification in certifications:

                st.write(
                    f"• {safe_string(certification)}"
                )

        else:

            st.write(
                "No certifications identified."
            )


        # ====================================================
        # EXPERIENCE EVIDENCE
        # ====================================================

        st.subheader("📌 Relevant Experience Evidence")

        evidence = detail_candidate.get(
            "relevant_experience_evidence",
            []
        )

        if evidence:

            for item in evidence:

                st.write(
                    f"• {safe_string(item)}"
                )

        else:

            st.write(
                "No specific experience evidence identified."
            )


        # ====================================================
        # STRENGTHS
        # ====================================================

        st.subheader("💪 Strengths")

        strengths = detail_candidate.get(
            "strengths",
            []
        )

        if strengths:

            for item in strengths:

                st.success(
                    safe_string(item)
                )

        else:

            st.write(
                "No specific strengths identified."
            )


        # ====================================================
        # GAPS
        # ====================================================

        st.subheader("⚠️ Potential Gaps")

        gaps = detail_candidate.get(
            "potential_gaps",
            []
        )

        if gaps:

            for item in gaps:

                st.warning(
                    safe_string(item)
                )

        else:

            st.success(
                "No significant gaps identified."
            )


        # ====================================================
        # SCORE BREAKDOWN
        # ====================================================

        st.subheader("⚖️ Score Breakdown")

        score_breakdown = pd.DataFrame({

            "Category": [
                "Experience",
                "Technical Skills",
                "Education",
                "Job Requirements",
                "Certifications"
            ],

            "Score": [

                detail_candidate.get(
                    "experience_score",
                    0
                ),

                detail_candidate.get(
                    "technical_score",
                    0
                ),

                detail_candidate.get(
                    "education_score",
                    0
                ),

                detail_candidate.get(
                    "requirements_score",
                    0
                ),

                detail_candidate.get(
                    "certification_score",
                    0
                )
            ],

            "Weight": [

                st.session_state.weight_experience,

                st.session_state.weight_technical,

                st.session_state.weight_education,

                st.session_state.weight_requirements,

                st.session_state.weight_certification
            ]
        })

        st.dataframe(
            score_breakdown,
            use_container_width=True,
            hide_index=True
        )


# ============================================================
# EXCEL EXPORT
# ============================================================

if st.session_state.candidates:

    st.header("📥 Export Results")

    ranking_export = []
    details_export = []
    keyword_export = []

    required_keywords = [

        normalize_keyword(x)

        for x in
        st.session_state.required_keywords.split(",")

        if normalize_keyword(x)
    ]

    for index, candidate in enumerate(
        st.session_state.candidates,
        start=1
    ):

        ranking_export.append({

            "Rank":
                index,

            "Candidate Name":
                candidate.get(
                    "candidate_name",
                    ""
                ),

            "File":
                candidate.get(
                    "file_name",
                    ""
                ),

            "Relevant Experience (Years)":
                candidate.get(
                    "years_relevant_experience",
                    0
                ),

            "Experience Score":
                candidate.get(
                    "experience_score",
                    0
                ),

            "Technical Score":
                candidate.get(
                    "technical_score",
                    0
                ),

            "Education Score":
                candidate.get(
                    "education_score",
                    0
                ),

            "Requirements Score":
                candidate.get(
                    "requirements_score",
                    0
                ),

            "Certification Score":
                candidate.get(
                    "certification_score",
                    0
                ),

            "Final Score":
                candidate.get(
                    "final_score",
                    0
                )
        })

        details_export.append({

            "Candidate Name":
                candidate.get(
                    "candidate_name",
                    ""
                ),

            "Education":
                safe_join(
                    candidate.get(
                        "education",
                        []
                    )
                ),

            "Previous Roles":
                safe_join(
                    candidate.get(
                        "previous_roles",
                        []
                    )
                ),

            "Technical Skills":
                safe_join(
                    candidate.get(
                        "technical_skills",
                        []
                    )
                ),

            "Software Tools":
                safe_join(
                    candidate.get(
                        "software_tools",
                        []
                    )
                ),

            "Certifications":
                safe_join(
                    candidate.get(
                        "certifications",
                        []
                    )
                ),

            "Experience Evidence":
                safe_join(
                    candidate.get(
                        "relevant_experience_evidence",
                        []
                    )
                ),

            "Strengths":
                safe_join(
                    candidate.get(
                        "strengths",
                        []
                    )
                ),

            "Potential Gaps":
                safe_join(
                    candidate.get(
                        "potential_gaps",
                        []
                    )
                )
        })

        keyword_export.append({

            "Candidate Name":
                candidate.get(
                    "candidate_name",
                    ""
                ),

            "Matched Keywords":
                safe_join(
                    candidate.get(
                        "matched_keywords",
                        []
                    )
                ),

            "Missing Keywords":
                safe_join(
                    candidate.get(
                        "missing_keywords",
                        []
                    )
                ),

            "Keyword Match %":
                keyword_score(
                    candidate.get(
                        "matched_keywords",
                        []
                    ),
                    required_keywords
                )
        })


    ranking_export_df = pd.DataFrame(
        ranking_export
    )

    details_export_df = pd.DataFrame(
        details_export
    )

    keyword_export_df = pd.DataFrame(
        keyword_export
    )

    weights_export_df = pd.DataFrame([

        {
            "Scoring Category": "Experience",
            "Weight":
                st.session_state.weight_experience
        },

        {
            "Scoring Category": "Technical Skills",
            "Weight":
                st.session_state.weight_technical
        },

        {
            "Scoring Category": "Education",
            "Weight":
                st.session_state.weight_education
        },

        {
            "Scoring Category": "Job Requirements",
            "Weight":
                st.session_state.weight_requirements
        },

        {
            "Scoring Category": "Certifications",
            "Weight":
                st.session_state.weight_certification
        }
    ])


    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        ranking_export_df.to_excel(
            writer,
            sheet_name="Candidate Ranking",
            index=False
        )

        details_export_df.to_excel(
            writer,
            sheet_name="Candidate Details",
            index=False
        )

        weights_export_df.to_excel(
            writer,
            sheet_name="Scoring Weights",
            index=False
        )

        keyword_export_df.to_excel(
            writer,
            sheet_name="Keyword Matching",
            index=False
        )

    output.seek(0)

    st.download_button(

        label="📥 Download Excel Report",

        data=output,

        file_name="HireTech_AI_Candidate_Ranking.xlsx",

        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),

        use_container_width=True
    )


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "HireTech AI • AI-assisted recruitment and candidate analysis"
)
